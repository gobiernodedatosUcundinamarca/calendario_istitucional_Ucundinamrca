"""
Lo que comparten las herramientas de datos: rutas, texto, catálogos de la app y el Excel fuente.

El Excel fuente (data/Calendario institucional.xlsx) es del que la app toma los datos: una hoja «Calendario»
con una fila por actividad y fecha. consolidar.py e importar-2026.py lo REEMPLAZAN (antes guardan una copia en
herramientas/salida/respaldos) y cargar.py genera desde él src/datos/actividades.ts y src/datos/responsables.ts.
"""
from __future__ import annotations

import datetime as dt
import re
import shutil
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# En el .exe (PyInstaller) las carpetas van junto al ejecutable, que se guarda en herramientas/.
HERRAMIENTAS = Path(sys.executable if getattr(sys, 'frozen', False) else __file__).resolve().parent
RAIZ = HERRAMIENTAS.parent
FUENTE = RAIZ / 'data' / 'Calendario institucional.xlsx'
HOJA_FUENTE = 'Calendario'
SALIDA = HERRAMIENTAS / 'salida'
RESPALDOS = SALIDA / 'respaldos'
CATALOGOS_TS = RAIZ / 'src' / 'datos' / 'catalogos.ts'
TODAS = 'Todas las sedes'
DIAS =['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']
MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre',
         'noviembre', 'diciembre']

# Columnas de la hoja «Calendario», en el orden del calendario institucional 2026. «Día» y «Mes» salen de
# «Fecha inicio»; las tres últimas dicen de dónde viene cada fila. La app usa las que lee cargar.py.
COLUMNAS_FUENTE = [
    ('sede', 'Unidad regional', 16), ('calendario', 'Calendario', 15), ('lider', 'Unidad líder', 30),
    ('categoria', 'Categoría', 24), ('subcategoria', 'Subcategoría', 30), ('actividad', 'Actividad a desarrollar', 50),
    ('lugar', 'Lugar de desarrollo', 28), ('hora_inicio', 'Hora inicio', 11), ('hora_fin', 'Hora final', 10),
    ('responsable', 'Responsable', 34), ('observaciones', 'Observaciones', 40), ('desde', 'Fecha inicio', 12),
    ('hasta', 'Fecha fin', 12), ('dia', 'Día', 11), ('mes', 'Mes', 11),
    ('repeticion', 'Repetición', 26), ('ajuste', 'Ajuste', 34), ('origen', 'Origen', 34),
]


# ── Texto ───────────────────────────────────────────────────────────────
def texto(v) -> str:
    if v is None:
        return ''
    s = ILLEGAL_CHARACTERS_RE.sub('', str(v).replace('_x000D_', ' '))  # _x000D_: salto de línea mal exportado
    return re.sub(r'\s+', ' ', s).strip()


def plano(v) -> str:
    """minúsculas, sin tildes ni signos, con letras y números separados."""
    s = unicodedata.normalize('NFD', texto(v)).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'([a-z])(\d)', r'\1 \2', s)
    s = re.sub(r'(\d)([a-z])', r'\1 \2', s)
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9]+', ' ', s)).strip()


# ── Catálogos de la app ─────────────────────────────────────────────────
def catalogos() -> dict[str, list[str]]:
    """Sedes, unidades líder, calendarios y categorías tal como están en src/datos/catalogos.ts, que manda sobre todo
    lo demás. Si el .exe se usa fuera del repositorio, toma la copia de catalogos.ts que lleva adentro."""
    ruta = CATALOGOS_TS if CATALOGOS_TS.exists() else Path(getattr(sys, '_MEIPASS', HERRAMIENTAS)) / CATALOGOS_TS.name
    ts = ruta.read_text(encoding='utf-8')
    literal = r"'((?:[^'\\]|\\.)*)'"

    def lista(nombre: str, abre: str, cierra: str, patron: str) -> list[str]:
        m = re.search(rf'export const {nombre} = {re.escape(abre)}(.*?){re.escape(cierra)}', ts, re.S)
        if not m:
            sys.exit(f'No encontré «{nombre}» en {CATALOGOS_TS}: ¿cambió su forma?')
        return [v.replace("\\'", "'").replace('\\\\', '\\') for v in re.findall(patron, m.group(1))]

    return {
        'sedes': lista('unidadesRegionales', '[', ']', literal),
        'lideres': lista('unidadesLider', '[', ']', literal),
        'calendarios': lista('calendarios', '[', ']', literal),
        'categorias': lista('categorias', '{', '}', literal + r'\s*:'),
    }


CALENDARIOS = catalogos()['calendarios']


# ── Fechas y horas del Excel fuente ─────────────────────────────────────
def a_fecha(v) -> dt.date | None:
    """Fecha de Excel o texto dd/mm/aaaa o aaaa-mm-dd (el Excel fuente se puede corregir a mano)."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    m = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', texto(v)) or re.fullmatch(r'(\d{4})-(\d{1,2})-(\d{1,2})', texto(v))
    if not m:
        return None
    a, b, c = map(int, m.groups())
    try:
        return dt.date(a, b, c) if a > 31 else dt.date(c, b, a)
    except ValueError:
        return None


def a_hora(v) -> str | None:
    """«HH:MM» desde una hora de Excel o un texto hh:mm."""
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return f'{v.hour:02d}:{v.minute:02d}'
    m = re.fullmatch(r'(\d{1,2}):(\d{2})(?::\d{2})?', texto(v))
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return f'{int(m.group(1)):02d}:{m.group(2)}'
    return None


# ── Hojas con formato ───────────────────────────────────────────────────
FUENTE_CELDA = Font(name='Arial', size=10)


def hoja_tabla(wb: Workbook, titulo: str, columnas: list[tuple[str, int]], filas: list[list], nota: str) -> None:
    """Una hoja con encabezado verde, filtros y anchos fijos. Si no hay filas, muestra la nota."""
    ws = wb.create_sheet(titulo)
    ws.append([n for n, _ in columnas])
    for fila in filas:
        ws.append(fila)
    for c, (_, ancho) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(c)].width = ancho
    for fila in ws.iter_rows():
        for celda in fila:
            celda.font = FUENTE_CELDA
            if isinstance(celda.value, str) and celda.value.startswith('='):
                celda.data_type = 's'  # un texto que empieza por «=» no debe volverse fórmula
            elif isinstance(celda.value, dt.date):
                celda.number_format = 'dd/mm/yyyy'
            elif isinstance(celda.value, dt.time):
                celda.number_format = 'hh:mm'
            elif isinstance(celda.value, float):
                celda.number_format = '0%'
    for celda in ws[1]:
        celda.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', start_color='007B3E')
        celda.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    if filas:
        tabla = Table(displayName=re.sub(r'[^A-Za-z0-9]', '', plano(titulo).title()) or 'Tabla',
                      ref=f'A1:{get_column_letter(len(columnas))}{len(filas) + 1}')
        tabla.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(tabla)
    else:
        ws.append([nota])
        ws['A2'].font = Font(name='Arial', size=10, italic=True)


def notas(ws, lineas: list[tuple[str, object]], titulo: str) -> None:
    """Hoja de texto: un título y pares «etiqueta, valor»."""
    ws.append([titulo])
    ws.append([])
    for etiqueta, valor in lineas:
        ws.append([etiqueta, valor])
    for fila in ws.iter_rows():
        for celda in fila:
            celda.font = FUENTE_CELDA
            celda.alignment = Alignment(wrap_text=True, vertical='top')
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color='007B3E')
    ws.column_dimensions['A'].width = 52
    ws.column_dimensions['B'].width = 90


def guardar(wb: Workbook, ruta: Path) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(ruta)
    except PermissionError:
        sys.exit(f'No se pudo escribir «{ruta}»: ciérrelo en Excel y vuelva a correr el script.')


# ── Excel fuente ────────────────────────────────────────────────────────
def filas_en_fuente(ruta: Path = FUENTE) -> int | None:
    """Cuántas filas tiene hoy el Excel fuente (None si no existe o no se puede leer)."""
    if not ruta.exists():
        return None
    try:
        ws = load_workbook(ruta, read_only=True, data_only=True)[HOJA_FUENTE]
        return sum(1 for fila in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, '') for v in fila))
    except Exception:  # noqa: BLE001 — si está dañado, igual se reemplaza (queda el respaldo)
        return None


def leer_filas_fuente(ruta: Path = FUENTE) -> list[dict]:
    """Filas de la hoja «Calendario» con las claves de COLUMNAS_FUENTE (más «fila»). Sale con un mensaje si el
    archivo no tiene esas columnas."""
    if not ruta.exists():
        sys.exit(f'No existe «{ruta}»: córralo después de consolidar.py o importar-2026.py.')
    ws = load_workbook(ruta, data_only=True)[HOJA_FUENTE]
    esperados = [t for _, t, _ in COLUMNAS_FUENTE]
    if [texto(c.value) for c in ws[1]][:len(esperados)] != esperados:
        sys.exit(f'La hoja «{HOJA_FUENTE}» de «{ruta.name}» no tiene las columnas esperadas: {", ".join(esperados)}')
    filas = []
    for numero, valores in enumerate(ws.iter_rows(min_row=2, max_col=len(esperados), values_only=True), start=2):
        if any(v not in (None, '') for v in valores):
            filas.append({'fila': numero, **{k: v for (k, _, _), v in zip(COLUMNAS_FUENTE, valores)}})
    return filas


def combinar_con_otros_anios(nuevas: list[dict], conservar: bool = True) -> tuple[list[dict], list[dict]]:
    """Las filas nuevas más las del Excel fuente actual que son de otros años (por «Fecha inicio»): consolidar
    2027 no borra 2026. Devuelve (todas, conservadas)."""
    conservadas = []
    if conservar and FUENTE.exists():
        anios = {f['desde'].year for f in nuevas if isinstance(f.get('desde'), dt.date)}
        for f in leer_filas_fuente(FUENTE):
            desde, hasta = a_fecha(f.get('desde')), a_fecha(f.get('hasta'))
            if desde and desde.year not in anios:
                conservadas.append({**f, 'desde': desde, 'hasta': hasta})
    todas = sorted(conservadas + nuevas, key=lambda f: (f['desde'], a_hora(f.get('hora_inicio')) or '', texto(f.get('actividad'))))
    return todas, conservadas


def con_dia_y_mes(f: dict) -> dict:
    desde = f.get('desde')
    if not isinstance(desde, dt.date):
        return f
    return {**f, 'dia': DIAS[desde.weekday()], 'mes': MESES[desde.month - 1]}


def escribir_fuente(ruta: Path, filas: list[dict], generado_por: str, origenes: list[str]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    hoja_tabla(wb, HOJA_FUENTE, [(t, a) for _, t, a in COLUMNAS_FUENTE],
               [[con_dia_y_mes(f).get(k) or None for k, _, _ in COLUMNAS_FUENTE] for f in filas], 'Sin actividades.')
    notas(wb.create_sheet('Acerca de'), [
        ('Qué es', 'El Excel del que la app web toma las actividades (hoja «Calendario», una fila por actividad y fecha).'),
        ('Generado por', f'{generado_por} el {dt.datetime.now():%d/%m/%Y %H:%M}'),
        ('A partir de', '\n'.join(origenes)),
        ('Actividades', len(filas)),
        ('Para corregir a mano', 'Edite la hoja «Calendario» y corra: python herramientas/cargar.py'),
        ('Cuidado', 'Volver a consolidar o importar REEMPLAZA las filas de los años que se cargan (las de otros años se '
                    'conservan) y se pierden los cambios hechos a mano en esos años (queda una copia en '
                    'herramientas/salida/respaldos). Corrija también en el formato del área.'),
        ('Columnas', f'Las del calendario institucional 2026. Unidad regional: una de las sedes o «{TODAS}». '
                     f'Calendario: {" o ".join(CALENDARIOS)}. Fechas dd/mm/aaaa y horas hh:mm (24 h). «Día» y «Mes» '
                     'salen de «Fecha inicio». Repetición, Ajuste y Origen dicen de dónde viene cada fila.'),
    ], 'Calendario institucional')
    guardar(wb, ruta)


def reemplazar_fuente(filas: list[dict], generado_por: str, origenes: list[str]) -> Path | None:
    """Escribe el Excel fuente con `filas` (ver combinar_con_otros_anios). Devuelve la copia del anterior (None si no había)."""
    respaldo = None
    if FUENTE.exists():
        respaldo = RESPALDOS / f'{FUENTE.stem} {dt.datetime.now():%Y-%m-%d %H%M%S}{FUENTE.suffix}'
        respaldo.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(FUENTE, respaldo)
    escribir_fuente(FUENTE, filas, generado_por, origenes)
    return respaldo
