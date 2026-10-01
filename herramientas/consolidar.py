"""
Consolida los formatos «Calendario Institucional - Áreas» que envía cada área y los carga en la app.

Uso (desde la raíz del repositorio):
  python herramientas/consolidar.py                   lee herramientas/formatos/ (incluidas subcarpetas)
  python herramientas/consolidar.py --anio 2027       además avisa de las fechas que no son de 2027
  python herramientas/consolidar.py --solo-revisar    no toca la app: deja el calendario en una vista previa
  python herramientas/consolidar.py --reemplazar-todo borra también los años que no vienen en los formatos
  python herramientas/consolidar.py --entrada <carpeta>

Salida:
  data/Calendario institucional.xlsx   Excel fuente de la app: se REEMPLAZAN los años que traen los formatos y se
                                       conservan los demás (copia del anterior en salida/respaldos)
  src/datos/actividades.ts, responsables.ts   los genera cargar.py desde el Excel fuente
  herramientas/salida/Revisión consolidación.xlsx   Resumen · Errores (filas que no entraron: se devuelven al
      área) · Avisos · Ajustes por festivo · Posibles duplicados · Reemplazos · Archivos

No confía en las fórmulas del formato: vuelve a revisar cada fila con las mismas reglas de «✔ Así quedó»,
porque pegar datos se salta las validaciones de Excel y Google Sheets o LibreOffice pueden no recalcularlas.

Reglas de consolidación:
  1. Solo se leen archivos .xlsx con las hojas y columnas del formato; los demás se rechazan con el motivo.
  2. Si un área (mismo «Área o dependencia» y «Sede principal») envía varios archivos, vale el más reciente
     (por «Fecha de envío» y, si empata o falta, por la fecha del archivo). Lo que solo estaba en el anterior
     queda en la hoja «Reemplazos» para confirmar que no era un envío adicional.
  3. Una fila con errores no entra al calendario: queda en «Errores» con todos sus problemas.
  4. Las repeticiones se expanden a una fila por fecha. Las sesiones que caen en festivo pasan al día hábil
     siguiente; una fecha puntual en festivo se deja igual y se avisa.
  5. Una fila idéntica a otra del mismo archivo se quita. Las parecidas (nombre casi igual, mismo día y sede,
     misma hora o sin hora) se conservan y se listan en «Posibles duplicados». Los cruces de horario se permiten.

Requiere Python 3.10+ y openpyxl.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import re
import sys
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from functools import cache, cached_property
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

import cargar
from comun import (CALENDARIOS, FUENTE, HERRAMIENTAS, RAIZ, SALIDA, TODAS, catalogos, combinar_con_otros_anios,
                   escribir_fuente, filas_en_fuente, guardar, hoja_tabla, notas, plano, reemplazar_fuente, texto)

if sys.stdout:  # en el .exe sin consola no hay salida estándar
    sys.stdout.reconfigure(encoding='utf-8')
warnings.filterwarnings('ignore', module='openpyxl')  # «Data Validation extension is not supported…»

HOY = dt.date.today()
REVISION = SALIDA / 'Revisión consolidación.xlsx'
VISTA_PREVIA = SALIDA / 'Calendario institucional (vista previa).xlsx'

# ── El formato ──────────────────────────────────────────────────────────
# Categorías, sedes y unidades líder salen de src/datos/catalogos.ts; la hoja «Listas» del formato debe coincidir.
CATALOGOS = catalogos()
CATEGORIAS = CATALOGOS['categorias']
SEDES = [*CATALOGOS['sedes'], TODAS]
LIDERES = CATALOGOS['lideres']
ALIAS_SEDES = {'bogota': 'Bogotá D.C.', 'bogota dc': 'Bogotá D.C.', 'fusa': 'Fusagasugá', 'zipa': 'Zipaquirá',
               'todas': TODAS, 'todas las unidades regionales': TODAS}
REPITE = ['No', 'Cada semana', 'Cada 15 días', 'Cada mes', 'Cada 2 meses', 'Cada 3 meses', 'Cada 6 meses',
          'En fechas específicas']
DIAS_REPITE = {'Cada semana': 7, 'Cada 15 días': 14}
MESES_REPITE = {'Cada mes': 1, 'Cada 2 meses': 2, 'Cada 3 meses': 3, 'Cada 6 meses': 6}

HOJA = 'Formato'
HOJAS_V2 = ('1. Datos del área', '2. Actividades')  # formato anterior, de dos hojas
DATOS_AREA = 'Datos del área'  # bloque de la parte de arriba de la hoja «Formato»
FILAS_AREA = range(1, 10)
# Etiqueta en la columna A (valor en B) y en la D (valor en F).
COLUMNAS_AREA = ((1, 2), (4, 6))
FILA_ENCABEZADO, PRIMERA_FILA, ULTIMA_FILA_FORMATO = 10, 11, 310
CAMPOS_AREA = {
    'area': 'Área o dependencia', 'lider': 'Unidad Líder', 'sede': 'Sede principal', 'calendario': 'Calendario principal',
    'quien': 'Nombre de quien diligencia', 'correo': 'Correo institucional', 'telefono': 'Extensión / teléfono',
    'envio': 'Fecha de envío', 'visto_bueno': 'Visto bueno del jefe del área (nombre y cargo)',
}
FECHAS = [f'fecha{i}' for i in range(1, 13)]
# Columnas A–Z de la tabla de «Formato». La M («✔ Así quedó») es una fórmula y no se lee.
COLUMNAS = ['actividad', 'categoria', 'subcategoria', 'responsable', 'calendario', 'sede', 'lugar', 'hora_inicio',
            'hora_fin', 'repite', 'desde', 'hasta', 'asi_quedo', *FECHAS, 'observaciones']
ENCABEZADOS = ['Actividad', 'Categoría', 'Subcategoría (opcional)', 'Responsable (vacío = el área)',
               'Calendario (vacío = el del área)', 'Sede (vacío = sede principal)', 'Lugar', 'Hora inicio (hh:mm)',
               'Hora fin (hh:mm)', '¿Se repite? (vacío = No)', 'Desde (dd/mm/aaaa)', 'Hasta (dd/mm/aaaa)',
               'Así quedó (automático)', *[f'Fecha {i}' for i in range(1, 13)], 'Observaciones']
# La primera versión del formato no pedía subcategoría, responsable ni calendario.
ENCABEZADOS_V1 = [e for e in ENCABEZADOS if e.split(' ')[0] not in ('Subcategoría', 'Responsable', 'Calendario')]
PLANTILLA = HERRAMIENTAS / 'plantilla' / 'Formato Calendario Institucional - Areas.xlsx'


def areas_sugeridas() -> dict[str, str]:
    """Las áreas de la hoja «Listas» (columna H) del formato vigente, para unificar cómo se escriben."""
    if not PLANTILLA.exists():
        return {}
    ws = load_workbook(PLANTILLA, read_only=True)['Listas']
    return {plano(v): texto(v) for (v,) in ws.iter_rows(min_row=2, min_col=8, max_col=8, values_only=True) if texto(v)}


AREAS = areas_sugeridas()


def nombre_area(valor) -> str:
    """El nombre tal como está en la lista si solo cambian mayúsculas, tildes o espacios; si no, como se escribió."""
    return AREAS.get(plano(valor), texto(valor))


DOMINIO_CORREO = 'ucundinamarca.edu.co'
UMBRAL_PARECIDO = 0.85


# ── Texto ───────────────────────────────────────────────────────────────
def vacio(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def opcion(valor, lista: list[str], alias: dict[str, str] | None = None) -> str | None:
    """El valor de la lista escrito igual, sin importar mayúsculas, tildes ni espacios."""
    p = plano(valor)
    return next((o for o in lista if plano(o) == p), None) or (alias or {}).get(p)


PALABRAS_VACIAS = {'de', 'del', 'la', 'las', 'los', 'el', 'y', 'e', 'o', 'u', 'a', 'al', 'en', 'con', 'por', 'para'}


def clave(nombre: str) -> str:
    """Nombre base para comparar: sin conectores ni años, en singular aproximado."""
    fichas = []
    for f in plano(nombre).split():
        if f in PALABRAS_VACIAS or re.fullmatch(r'(19|20)\d\d', f):
            continue
        if len(f) > 4 and f.endswith('ones'):
            f = f[:-2]
        elif len(f) > 4 and f.endswith('s') and f[-2] in 'aeiou':
            f = f[:-1]
        fichas.append(f)
    return ' '.join(fichas)


@cache
def parecido(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    corto, largo = sorted((a, b), key=len)
    if len(corto) >= 8 and f' {corto} ' in f' {largo} ':
        return 0.92
    return SequenceMatcher(None, a, b).ratio()


# ── Fechas y horas ──────────────────────────────────────────────────────
EPOCA_EXCEL = dt.date(1899, 12, 30)


def leer_fecha(v, campo: str) -> tuple[dt.date, str | None]:
    """(fecha, aviso). ValueError si no es una fecha."""
    error = f'«{campo}» no es una fecha válida: use dd/mm/aaaa'
    if isinstance(v, dt.datetime):
        aviso = f'«{campo}» traía también una hora: se tomó solo la fecha' if v.time() != dt.time(0) else None
        return v.date(), aviso
    if isinstance(v, dt.date):
        return v, None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 36526 <= v < 73051:  # 2000-01-01 a 2099-12-31 como número de serie de Excel
            return EPOCA_EXCEL + dt.timedelta(days=int(v)), f'«{campo}» venía como número ({v:g}): se leyó como fecha'
        raise ValueError(error)
    s = texto(v)
    for patron, orden in ((r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})', 'dma'), (r'(\d{4})-(\d{1,2})-(\d{1,2})', 'amd')):
        m = re.fullmatch(patron, s)
        if m:
            partes = dict(zip(orden, map(int, m.groups())))
            partes['a'] += 2000 if partes['a'] < 100 else 0
            try:
                fecha = dt.date(partes['a'], partes['m'], partes['d'])
            except ValueError:
                raise ValueError(error) from None
            return fecha, f'«{campo}» venía como texto («{s}»): se leyó como {fecha:%d/%m/%Y}'
    raise ValueError(error)


def leer_hora(v, campo: str) -> tuple[dt.time, str | None]:
    """(hora, aviso). ValueError si no es una hora."""
    error = f'«{campo}» no es una hora válida: use hh:mm en 24 horas (ej. 14:30)'
    aviso = None
    if isinstance(v, dt.datetime):
        if v.date() not in (dt.date(1899, 12, 30), dt.date(1899, 12, 31), dt.date(1900, 1, 1)):
            aviso = f'«{campo}» traía también una fecha: se tomó solo la hora'
        v = v.time()
    if isinstance(v, dt.time):
        return v.replace(second=0, microsecond=0), aviso
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if not 0 <= v < 1:
            raise ValueError(error)
        minutos = round(v * 24 * 60)
        if minutos >= 24 * 60:
            raise ValueError(error)
        return dt.time(minutos // 60, minutos % 60), f'«{campo}» venía como número: se leyó como hora'
    s = plano(v).replace(' ', '')  # «2:30 p. m.» → «230pm», «14h30» → «14h30», «14:30:00» → «143000»
    m = re.fullmatch(r'(\d{1,2})(?:h?(\d{2})(?:\d{2})?)?h?(am|pm|m)?', s)
    if not m:
        raise ValueError(error)
    h, mi, sufijo = int(m.group(1)), int(m.group(2) or 0), m.group(3)
    if sufijo == 'pm' and h < 12:
        h += 12
    elif sufijo == 'am' and h == 12:
        h = 0
    if h > 23 or mi > 59:
        raise ValueError(error)
    return dt.time(h, mi), f'«{campo}» venía como texto («{texto(v)}»): se leyó como {h:02d}:{mi:02d}'


def pascua(anio: int) -> dt.date:
    """Domingo de Pascua (algoritmo gregoriano anónimo)."""
    a, b, c = anio % 19, anio // 100, anio % 100
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes, dia = divmod(h + l - 7 * m + 114, 31)
    return dt.date(anio, mes, dia + 1)


@cache
def festivos(anio: int) -> dict[dt.date, str]:
    """Festivos de Colombia (Ley 51 de 1983): fijos, trasladables al lunes siguiente y los que dependen de Pascua."""
    def lunes(d: dt.date) -> dt.date:
        return d + dt.timedelta(days=-d.weekday() % 7)

    p = pascua(anio)
    dias = [
        (dt.date(anio, 1, 1), 'Año Nuevo'), (lunes(dt.date(anio, 1, 6)), 'Reyes Magos'),
        (lunes(dt.date(anio, 3, 19)), 'San José'), (p - dt.timedelta(days=3), 'Jueves Santo'),
        (p - dt.timedelta(days=2), 'Viernes Santo'), (dt.date(anio, 5, 1), 'Día del Trabajo'),
        (p + dt.timedelta(days=43), 'Ascensión del Señor'), (p + dt.timedelta(days=64), 'Corpus Christi'),
        (p + dt.timedelta(days=71), 'Sagrado Corazón'), (lunes(dt.date(anio, 6, 29)), 'San Pedro y San Pablo'),
        (dt.date(anio, 7, 20), 'Día de la Independencia'), (dt.date(anio, 8, 7), 'Batalla de Boyacá'),
        (lunes(dt.date(anio, 8, 15)), 'Asunción de la Virgen'), (lunes(dt.date(anio, 10, 12)), 'Día de la Raza'),
        (lunes(dt.date(anio, 11, 1)), 'Todos los Santos'), (lunes(dt.date(anio, 11, 11)), 'Independencia de Cartagena'),
        (dt.date(anio, 12, 8), 'Inmaculada Concepción'), (dt.date(anio, 12, 25), 'Navidad'),
    ]
    salida: dict[dt.date, str] = {}
    for d, nombre in dias:
        salida[d] = f'{salida[d]} y {nombre}' if d in salida else nombre
    return salida


def festivo(d: dt.date) -> str | None:
    return festivos(d.year).get(d)


def dia_habil_siguiente(d: dt.date) -> dt.date:
    d += dt.timedelta(days=1)
    while d.weekday() >= 5 or festivo(d):
        d += dt.timedelta(days=1)
    return d


def fechas_de_repeticion(modo: str, desde: dt.date, hasta: dt.date) -> list[dt.date]:
    """Las mismas fechas que calcula el formato (columna «N° fechas»)."""
    if modo in DIAS_REPITE:
        paso = DIAS_REPITE[modo]
        return [desde + dt.timedelta(days=n) for n in range(0, (hasta - desde).days + 1, paso)]
    # Mensual: el mismo día de la semana y la misma semana del mes que «Desde» (la 5.ª = la última del mes).
    cada, semana, dia_semana = MESES_REPITE[modo], (desde.day - 1) // 7 + 1, desde.weekday()
    salida, n = [], 0
    while True:
        anio, mes = divmod(desde.month - 1 + n, 12)
        primero = dt.date(desde.year + anio, mes + 1, 1)
        if primero > hasta:
            return salida
        primera = primero + dt.timedelta(days=(dia_semana - primero.weekday()) % 7)
        if semana < 5:
            d = primera + dt.timedelta(weeks=semana - 1)
        else:
            d = primera
            while (d + dt.timedelta(weeks=1)).month == primero.month:
                d += dt.timedelta(weeks=1)
        if desde <= d <= hasta:
            salida.append(d)
        n += cada


def dmy(d: dt.date) -> str:
    return f'{d:%d/%m/%Y}'


# ── Modelo ──────────────────────────────────────────────────────────────
class Rechazo(Exception):
    """El archivo completo no se puede consolidar."""


@dataclass
class Ocurrencia:
    inicio: dt.date
    fin: dt.date
    original: dt.date | None = None  # fecha antes de moverla por festivo


@dataclass(eq=False)
class Fila:
    archivo: Archivo
    numero: int
    actividad: str
    categoria: str
    subcategoria: str
    responsable: str
    calendario: str
    sede: str
    lugar: str
    hora_inicio: dt.time | None
    hora_fin: dt.time | None
    modo: str
    observaciones: str
    ocurrencias: list[Ocurrencia]

    @cached_property
    def clave(self) -> str:
        return clave(self.actividad)


@dataclass(eq=False)
class Archivo:
    ruta: Path
    estado: str = 'Aceptado'
    motivo: str = ''
    datos: dict = field(default_factory=dict)
    huella: str = ''
    modificado: dt.datetime | None = None
    escritas: int = 0
    filas: list[Fila] = field(default_factory=list)
    errores: list[tuple[int | str, str, str]] = field(default_factory=list)  # (fila, actividad, problemas)
    avisos: list[tuple[int | str, str, str]] = field(default_factory=list)

    @property
    def area(self) -> str:
        return texto(self.datos.get('area'))

    def rechazar(self, motivo: str) -> None:
        self.estado, self.motivo = 'Rechazado', motivo


# ── Lectura de un formato ───────────────────────────────────────────────
def leer_archivo(ruta: Path, anio: int | None) -> Archivo:
    a = Archivo(ruta)
    a.huella = hashlib.sha256(ruta.read_bytes()).hexdigest()
    a.modificado = dt.datetime.fromtimestamp(ruta.stat().st_mtime)
    try:
        if ruta.suffix.lower() != '.xlsx':
            raise Rechazo('No es .xlsx: ábralo en Excel y guárdelo como «Libro de Excel (.xlsx)»')
        try:
            wb = load_workbook(ruta, data_only=True)
        except Exception as e:  # noqa: BLE001 — cualquier archivo ilegible se reporta igual
            raise Rechazo(f'No se pudo abrir (dañado o protegido con contraseña): {type(e).__name__}') from None
        hojas = {plano(n): wb[n] for n in wb.sheetnames}
        hoja = hojas.get(plano(HOJA))
        if hoja is None:
            if all(plano(n) in hojas for n in HOJAS_V2):
                raise Rechazo('Usa la versión anterior del formato, de dos hojas: envíele el formato vigente (una sola hoja, '
                              '«Formato») y pídale que pase allí sus actividades')
            raise Rechazo(f'No es el formato: falta la hoja «{HOJA}» (tiene: {", ".join(wb.sheetnames)})')
        revisar_encabezados(hoja)
        leer_datos_area(a, hoja)
        revisar_version(a, hojas)
        leer_actividades(a, hoja, anio)
    except Rechazo as r:
        a.rechazar(str(r))
    return a


def leer_datos_area(a: Archivo, hoja) -> None:
    celdas = {plano(hoja.cell(r, col).value): (r, valor) for r in FILAS_AREA for col, valor in COLUMNAS_AREA
              if hoja.cell(r, col).value}
    for campo, etiqueta in CAMPOS_AREA.items():
        if plano(etiqueta) not in celdas:
            raise Rechazo(f'Se modificó la hoja «{HOJA}»: no está el campo «{etiqueta}» de los {DATOS_AREA.lower()}')
        r, col = celdas[plano(etiqueta)]
        a.datos[campo] = hoja.cell(r, col).value

    def aviso(mensaje: str) -> None:
        a.avisos.append(('—', '', mensaje))

    if not a.area:
        raise Rechazo(f'Falta «Área o dependencia» en «{DATOS_AREA}»: no se sabe de quién son las actividades')
    a.datos['area'] = nombre_area(a.datos['area'])
    lider = opcion(a.datos['lider'], LIDERES)
    if not lider:
        raise Rechazo(f'Falta la «Unidad Líder» en «{DATOS_AREA}» o no es una de la lista')
    a.datos['lider'] = lider
    sede = opcion(a.datos['sede'], SEDES, ALIAS_SEDES)
    if not vacio(a.datos['sede']) and not sede:
        aviso(f'La «Sede principal» «{texto(a.datos["sede"])}» no está en la lista: las filas sin sede quedarán con error')
    a.datos['sede'] = sede
    calendario = opcion(a.datos['calendario'], CALENDARIOS)
    if not vacio(a.datos['calendario']) and not calendario:
        aviso(f'El «Calendario principal» «{texto(a.datos["calendario"])}» no es Académico ni Administrativo: '
              'las filas sin calendario quedarán con error')
    a.datos['calendario'] = calendario
    envio = a.datos['envio']
    try:
        a.datos['envio'] = None if vacio(envio) else leer_fecha(envio, 'Fecha de envío')[0]
    except ValueError:
        a.datos['envio'] = None
        aviso(f'La «Fecha de envío» «{texto(envio)}» no es una fecha: se usa la fecha del archivo')
    for campo in ('quien', 'correo', 'visto_bueno'):
        if vacio(a.datos[campo]):
            aviso(f'Falta «{CAMPOS_AREA[campo]}» en «{DATOS_AREA}»')
    correo = texto(a.datos['correo'])
    if correo and not re.fullmatch(rf'[^@\s]+@{re.escape(DOMINIO_CORREO)}', correo, re.IGNORECASE):
        aviso(f'El correo «{correo}» no es institucional (@{DOMINIO_CORREO})')


def revisar_version(a: Archivo, hojas: dict) -> None:
    """Un área puede estar usando una copia vieja del formato (con otras sedes, categorías o unidades)."""
    listas = hojas.get('listas')
    if listas is None:
        a.avisos.append(('—', '', 'No tiene la hoja «Listas»: puede ser otra versión del formato'))
    else:
        esperadas = {1: CATEGORIAS, 2: SEDES, 3: REPITE, 4: LIDERES, 7: CALENDARIOS}
        distintas = []
        for col, esperada in esperadas.items():
            valores = [texto(listas.cell(r, col).value) for r in range(2, 60) if texto(listas.cell(r, col).value)]
            if set(valores) != set(esperada):
                distintas.append(texto(listas.cell(1, col).value) or f'columna {get_column_letter(col)}')
        if distintas:
            a.avisos.append(('—', '', f'Usa otra versión del formato (cambian las listas de: {", ".join(distintas)}). Envíele el formato vigente'))


def revisar_encabezados(hoja) -> None:
    encontrados = [plano(hoja.cell(FILA_ENCABEZADO, c).value) for c in range(1, len(ENCABEZADOS) + 1)]
    if encontrados[:len(ENCABEZADOS_V1)] == [plano(e) for e in ENCABEZADOS_V1]:
        raise Rechazo('Usa la versión anterior del formato, sin «Subcategoría», «Responsable» ni «Calendario»: '
                      'envíele el formato vigente y pídale que pase allí sus actividades')
    for c, (esperado, encontrado) in enumerate(zip(ENCABEZADOS, encontrados), start=1):
        if plano(esperado) != encontrado:
            raise Rechazo(f'Se modificaron las columnas de «{HOJA}»: la columna {get_column_letter(c)} '
                          f'dice «{texto(hoja.cell(FILA_ENCABEZADO, c).value)}» y debería decir «{esperado}»')


def leer_actividades(a: Archivo, hoja, anio: int | None) -> None:
    combinadas = [str(r) for r in hoja.merged_cells.ranges if r.max_row >= PRIMERA_FILA and r.min_col <= len(COLUMNAS)]
    if combinadas:
        a.avisos.append(('—', '', f'Hay celdas combinadas en la tabla de «{HOJA}» ({", ".join(combinadas[:5])}): '
                                  'solo se lee la primera celda de cada grupo'))
    vacias_seguidas = 0
    for numero, valores in enumerate(hoja.iter_rows(min_row=PRIMERA_FILA, max_col=len(COLUMNAS), values_only=True),
                                     start=PRIMERA_FILA):
        c = dict(zip(COLUMNAS, valores))
        if all(vacio(v) for k, v in c.items() if k not in ('asi_quedo', 'observaciones')):
            if not vacio(c['observaciones']):
                a.avisos.append((numero, '', 'La fila solo tiene «Observaciones»: no se tomó'))
            vacias_seguidas += 1
            if numero > ULTIMA_FILA_FORMATO and vacias_seguidas > 200:
                break
            continue
        vacias_seguidas = 0
        a.escritas += 1
        nombre = texto(c['actividad'])
        fila, errores, avisos = revisar_fila(a, numero, c, anio)
        if numero > ULTIMA_FILA_FORMATO:
            avisos.insert(0, f'Está debajo de la fila {ULTIMA_FILA_FORMATO}, fuera del formato: Excel no la revisó')
        a.avisos.extend((numero, nombre, m) for m in avisos)
        if errores:
            a.errores.append((numero, nombre, ' · '.join(errores)))
        else:
            a.filas.append(fila)


def revisar_fila(a: Archivo, numero: int, c: dict, anio: int | None) -> tuple[Fila | None, list[str], list[str]]:
    """Las reglas de la columna «✔ Así quedó» del formato, todas a la vez (el formato muestra solo la primera)."""
    errores, avisos = [], []

    def leer(funcion, valor, campo):
        if vacio(valor):
            return None, None
        try:
            v, aviso = funcion(valor, campo)
        except ValueError as e:
            return None, str(e)
        if aviso:
            avisos.append(aviso)
        return v, None

    nombre = texto(c['actividad'])
    if not nombre:
        errores.append('Falta el nombre de la actividad')

    categoria = opcion(c['categoria'], CATEGORIAS)
    if vacio(c['categoria']):
        errores.append('Falta la categoría')
    elif not categoria:
        errores.append(f'La categoría «{texto(c["categoria"])}» no está en la lista')

    if vacio(c['sede']):
        sede = a.datos['sede']
        if not sede:
            errores.append(f'Falta la sede: elíjala en la fila o en «Sede principal» de «{DATOS_AREA}»')
    else:
        sede = opcion(c['sede'], SEDES, ALIAS_SEDES)
        if not sede:
            errores.append(f'La sede «{texto(c["sede"])}» no está en la lista')
        elif sede != texto(c['sede']):
            avisos.append(f'Sede «{texto(c["sede"])}» leída como «{sede}»')

    if vacio(c['calendario']):
        calendario = a.datos['calendario']
        if not calendario:
            errores.append(f'Falta el calendario (Académico o Administrativo): elíjalo en la fila o en '
                           f'«Calendario principal» de «{DATOS_AREA}»')
    else:
        calendario = opcion(c['calendario'], CALENDARIOS)
        if not calendario:
            errores.append(f'El calendario «{texto(c["calendario"])}» no es Académico ni Administrativo')

    responsable = nombre_area(c['responsable']) or a.area
    subcategoria = texto(c['subcategoria'])

    lugar = texto(c['lugar'])
    if not lugar:
        errores.append('Falta el lugar')

    hora_inicio, error_hi = leer(leer_hora, c['hora_inicio'], 'Hora inicio')
    hora_fin, error_hf = leer(leer_hora, c['hora_fin'], 'Hora fin')
    errores += [e for e in (error_hi, error_hf) if e]
    if hora_inicio and hora_fin:
        if hora_fin < hora_inicio:
            errores.append('La hora fin es antes de la hora inicio')
        elif hora_fin == hora_inicio:
            avisos.append('La hora fin es igual a la hora inicio')
    elif hora_fin and vacio(c['hora_inicio']):
        avisos.append('Tiene hora fin pero no hora inicio')

    repite = ''
    if not vacio(c['repite']):
        repite = opcion(c['repite'], REPITE)
        if not repite:
            errores.append(f'«¿Se repite?»: «{texto(c["repite"])}» no es una opción de la lista')

    hay_sueltas = any(not vacio(c[f]) for f in FECHAS)
    desde, error_desde = leer(leer_fecha, c['desde'], 'Desde')
    hasta, error_hasta = leer(leer_fecha, c['hasta'], 'Hasta')
    fechas: list[dt.date] = []
    modo = None
    if repite == 'En fechas específicas' or (repite in ('', 'No') and vacio(c['desde']) and hay_sueltas):
        modo = 'Fechas específicas'
        for i, f in enumerate(FECHAS, start=1):
            valor, error = leer(leer_fecha, c[f], f'Fecha {i}')
            if error:
                errores.append(error)
            elif valor:
                fechas.append(valor)
        if not hay_sueltas:
            errores.append('Escriba las fechas en Fecha 1, Fecha 2…')
        if len(set(fechas)) < len(fechas):
            avisos.append('Hay fechas repetidas en Fecha 1…12: se tomó cada una una vez')
        if not vacio(c['desde']) or not vacio(c['hasta']):
            avisos.append('Con «En fechas específicas» no se usan «Desde» ni «Hasta»: se tomaron Fecha 1…12')
        fechas = sorted(set(fechas))
    elif repite is not None:
        modo = repite if repite not in ('', 'No') else (
            'Un día' if vacio(c['hasta']) or (desde and hasta and desde == hasta) else 'Varios días')
        if hay_sueltas:
            errores.append('Escribió fechas en Fecha 1…12: elija «En fechas específicas» o bórrelas')
        if vacio(c['desde']):
            errores.append('Falta la fecha en «Desde»')
        errores += [e for e in (error_desde, error_hasta) if e]
        if desde and hasta and hasta < desde:
            errores.append('«Hasta» es antes de «Desde»')
        if (modo in DIAS_REPITE or modo in MESES_REPITE) and vacio(c['hasta']):
            errores.append('Falta «Hasta»: la última fecha en que se repite')
        if desde and hasta and (hasta - desde).days > 730:
            errores.append('El periodo no puede pasar de 2 años')

    primera = fechas[0] if fechas else desde
    if primera and primera.year < HOY.year:
        errores.append(f'La fecha {dmy(primera)} es de un año que ya pasó: revise el año')

    if errores:
        return None, errores, avisos

    ocurrencias = calcular_ocurrencias(modo, desde, hasta, fechas, avisos)
    pasadas = sum(o.fin < HOY for o in ocurrencias)
    if pasadas:
        avisos.append(f'{pasadas} de {len(ocurrencias)} fechas ya pasaron: ¿el año está bien?')
    if anio:
        fuera = sum(o.fin.year < anio or o.inicio.year > anio for o in ocurrencias)
        if fuera:
            avisos.append(f'{fuera} de {len(ocurrencias)} fechas no son de {anio}')
    fila = Fila(a, numero, nombre, categoria, subcategoria, responsable, calendario, sede, lugar, hora_inicio, hora_fin,
                modo, texto(c['observaciones']), ocurrencias)
    return fila, errores, avisos


def calcular_ocurrencias(modo: str, desde, hasta, fechas: list[dt.date], avisos: list[str]) -> list[Ocurrencia]:
    if modo == 'Varios días':
        return [Ocurrencia(desde, hasta)]
    if modo in ('Un día', 'Fechas específicas'):
        dias = fechas if modo == 'Fechas específicas' else [desde]
        en_festivo = [f'{dmy(d)} ({festivo(d)})' for d in dias if festivo(d)]
        if en_festivo:
            avisos.append(f'Cae en festivo: {", ".join(en_festivo)}. Se dejó igual: confirme la fecha')
        return [Ocurrencia(d, d) for d in dias]
    salida, vistas = [], set()
    for d in fechas_de_repeticion(modo, desde, hasta):
        nueva = dia_habil_siguiente(d) if festivo(d) else d
        if nueva in vistas:
            continue
        vistas.add(nueva)
        salida.append(Ocurrencia(nueva, nueva, d if nueva != d else None))
    return salida


# ── Consolidación ───────────────────────────────────────────────────────
def resolver_reenvios(archivos: list[Archivo]) -> list[tuple[Archivo, Archivo, Fila]]:
    """Archivos idénticos o del mismo área: vale el más reciente. Devuelve lo que solo estaba en un anterior."""
    por_huella: dict[str, Archivo] = {}
    for a in archivos:
        if a.estado != 'Aceptado':
            continue
        if a.huella in por_huella:
            a.estado, a.motivo = 'Repetido', f'Idéntico a «{por_huella[a.huella].ruta.name}»'
        else:
            por_huella[a.huella] = a
    grupos: dict[tuple, list[Archivo]] = defaultdict(list)
    for a in archivos:
        if a.estado == 'Aceptado':
            grupos[(plano(a.area), a.datos['sede'])].append(a)
    solo_en_anterior = []
    for grupo in grupos.values():
        if len(grupo) < 2:
            continue
        grupo.sort(key=lambda a: (a.datos['envio'] or dt.date.min, a.modificado))
        nuevo = grupo[-1]
        claves_nuevo = {f.clave for f in nuevo.filas} | {clave(n) for _, n, _ in nuevo.errores}
        for viejo in grupo[:-1]:
            criterio = ('fecha de envío' if viejo.datos['envio'] and nuevo.datos['envio']
                        and viejo.datos['envio'] != nuevo.datos['envio'] else 'fecha del archivo')
            viejo.estado = 'Reemplazado'
            viejo.motivo = f'Hay una versión más reciente del área ({criterio}): «{nuevo.ruta.name}»'
            for f in viejo.filas:
                if not any(parecido(f.clave, c) >= UMBRAL_PARECIDO for c in claves_nuevo):
                    solo_en_anterior.append((viejo, nuevo, f))
    return solo_en_anterior


def quitar_repetidas(filas: list[Fila]) -> tuple[list[Fila], list[tuple]]:
    """Filas idénticas dentro de un mismo archivo (copiar y pegar dos veces)."""
    vistas: dict[tuple, Fila] = {}
    quedan, quitadas = [], []
    for f in filas:
        llave = (f.archivo.ruta, plano(f.actividad), f.categoria, plano(f.subcategoria), plano(f.responsable), f.calendario,
                 f.sede, plano(f.lugar), f.hora_inicio, f.hora_fin, tuple((o.inicio, o.fin) for o in f.ocurrencias))
        if llave in vistas:
            primera = vistas[llave]
            quitadas.append((primera, f, 1.0, len(f.ocurrencias), f.ocurrencias[0].inicio,
                             'Se quitó la fila B: es idéntica a la A en el mismo archivo'))
        else:
            vistas[llave] = f
            quedan.append(f)
    return quedan, quitadas


def posibles_duplicados(filas: list[Fila]) -> list[tuple]:
    """Nombre casi igual, el mismo día y sede, y la misma hora (o sin hora): se conservan, para revisar."""
    por_dia: dict[dt.date, list[Fila]] = defaultdict(list)
    for f in filas:
        for o in f.ocurrencias:
            por_dia[o.inicio].append(f)
    pares: dict[tuple[int, int], list] = {}
    for dia, del_dia in sorted(por_dia.items()):
        for i, x in enumerate(del_dia):
            for y in del_dia[i + 1:]:
                if x is y or not (x.sede == y.sede or TODAS in (x.sede, y.sede)):
                    continue
                if x.hora_inicio and y.hora_inicio and x.hora_inicio != y.hora_inicio:
                    continue
                similitud = parecido(x.clave, y.clave)
                if similitud < UMBRAL_PARECIDO:
                    continue
                llave = (id(x), id(y))
                if llave in pares:
                    pares[llave][3] += 1
                else:
                    pares[llave] = [x, y, similitud, 1, dia]
    salida = []
    for x, y, similitud, comunes, primera in pares.values():
        que = ('Se conservaron las dos: revisar (mismo archivo)' if x.archivo is y.archivo
               else 'Se conservaron las dos: revisar (¿la reportaron dos áreas?)')
        salida.append((x, y, similitud, comunes, primera, que))
    return salida


# ── Salida ──────────────────────────────────────────────────────────────
def filas_fuente(filas: list[Fila]) -> list[dict]:
    """Una fila por fecha, en las columnas del Excel fuente (comun.COLUMNAS_FUENTE)."""
    salida = []
    for f in filas:
        total = len(f.ocurrencias)
        for n, o in enumerate(f.ocurrencias, start=1):
            salida.append(dict(
                actividad=f.actividad, categoria=f.categoria, desde=o.inicio, hasta=o.fin, hora_inicio=f.hora_inicio,
                hora_fin=f.hora_fin, sede=f.sede, lugar=f.lugar, responsable=f.responsable, calendario=f.calendario,
                subcategoria=f.subcategoria, lider=f.archivo.datos['lider'], observaciones=f.observaciones,
                repeticion=f'{f.modo} · {n} de {total}' if total > 1 else f.modo,
                ajuste=f'Movida del {dmy(o.original)} ({festivo(o.original)})' if o.original else None,
                origen=f'{f.archivo.ruta.name} · fila {f.numero}'))
    salida.sort(key=lambda r: (r['desde'], r['hora_inicio'] or dt.time(0), r['actividad']))
    return salida


def escribir_revision(entrada: Path, archivos: list[Archivo], filas: list[Fila], fechas: int, quitadas: list[tuple],
                      duplicados: list[tuple], solo_en_anterior: list[tuple], anio: int | None,
                      destino: list[tuple[str, str]]) -> dict:
    finales = [a for a in archivos if a.estado == 'Aceptado']
    rechazados = [a for a in archivos if a.estado == 'Rechazado']

    def contacto(a: Archivo) -> list:
        return [a.ruta.name, a.area, texto(a.datos.get('quien')), texto(a.datos.get('correo'))]

    errores =[[*contacto(a), '—', '', f'Archivo rechazado: {a.motivo}'] for a in rechazados]
    errores += [[*contacto(a), n, nombre, problemas] for a in finales for n, nombre, problemas in a.errores]
    avisos = [[a.ruta.name, a.area, n, nombre, m] for a in finales for n, nombre, m in a.avisos]
    ajustes = [[f.archivo.ruta.name, f.archivo.area, f.numero, f.actividad, f.modo, o.original,
                festivo(o.original), o.inicio] for f in filas for o in f.ocurrencias if o.original]
    pares = [[x.actividad, x.archivo.area, x.archivo.ruta.name, x.numero, y.actividad, y.archivo.area,
              y.archivo.ruta.name, y.numero, similitud, comunes, primera, x.sede if x.sede == y.sede else f'{x.sede} / {y.sede}', que]
             for x, y, similitud, comunes, primera, que in quitadas + duplicados]
    reemplazos = [[viejo.ruta.name, nuevo.ruta.name, viejo.area, f.numero, f.actividad, f.ocurrencias[0].inicio]
                  for viejo, nuevo, f in solo_en_anterior]
    lista_archivos = [[a.ruta.name, a.estado, a.motivo, a.area, a.datos.get('lider') or '', a.datos.get('sede') or '',
                       texto(a.datos.get('quien')), texto(a.datos.get('correo')), texto(a.datos.get('telefono')),
                       a.datos.get('envio'), texto(a.datos.get('visto_bueno')), a.escritas, len(a.errores),
                       sum(len(f.ocurrencias) for f in a.filas) if a.estado == 'Aceptado' else 0,
                       a.modificado.date() if a.modificado else None, a.huella[:12]]
                      for a in archivos]
    estados = Counter(a.estado for a in archivos)
    resumen = {
        'Archivos encontrados': len(archivos),
        'Aceptados': estados['Aceptado'],
        'Rechazados (ver «Errores»)': estados['Rechazado'],
        'Reemplazados por una versión más reciente': estados['Reemplazado'],
        'Repetidos (idénticos a otro)': estados['Repetido'],
        'Ignorados (no son hojas de cálculo)': estados['Ignorado'],
        'Filas escritas en los aceptados': sum(a.escritas for a in finales),
        'Filas con errores (no entraron)': sum(len(a.errores) for a in finales),
        'Filas repetidas quitadas': len(quitadas),
        'Actividades consolidadas': len(filas),
        'Fechas en el calendario': fechas,
        'Sesiones movidas por festivo': len(ajustes),
        'Posibles duplicados (revisar)': len(duplicados),
        'Actividades que solo estaban en un archivo reemplazado': len(reemplazos),
        'Avisos': len(avisos),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    notas(ws, [
        ('Generado', f'{dt.datetime.now():%d/%m/%Y %H:%M} desde «{entrada}»' + (f' · año {anio}' if anio else '')),
        *destino, ('', ''), *resumen.items(), ('', ''),
        ('Qué revisar', '1. «Errores»: devuelva cada fila al área (el correo está en la misma hoja). No entraron al calendario.\n'
                        '2. «Posibles duplicados» y «Reemplazos»: decida a mano; el script no borra nada de eso.\n'
                        '3. «Ajustes por festivo» y «Avisos»: confirme las fechas movidas o leídas de texto.\n'
                        'Si corrige algo, hágalo en el formato del área y vuelva a consolidar.'),
    ], 'Consolidación del calendario institucional')

    hoja_tabla(wb, 'Errores', [('Archivo', 30), ('Área', 30), ('Quien diligencia', 24), ('Correo', 30), ('Fila', 6),
                               ('Actividad', 44), ('Qué corregir', 90)], errores, 'Sin errores.')
    hoja_tabla(wb, 'Avisos', [('Archivo', 30), ('Área', 30), ('Fila', 6), ('Actividad', 44), ('Aviso', 90)],
               avisos, 'Sin avisos.')
    hoja_tabla(wb, 'Ajustes por festivo', [('Archivo', 30), ('Área', 30), ('Fila', 6), ('Actividad', 44),
                                           ('Repetición', 16), ('Fecha original', 14), ('Festivo', 26),
                                           ('Nueva fecha', 14)], ajustes, 'Ninguna sesión cayó en festivo.')
    hoja_tabla(wb, 'Posibles duplicados', [('Actividad A', 40), ('Área A', 26), ('Archivo A', 24), ('Fila A', 7),
                                           ('Actividad B', 40), ('Área B', 26), ('Archivo B', 24), ('Fila B', 7),
                                           ('Parecido', 9), ('Fechas en común', 9), ('Primera fecha en común', 12),
                                           ('Sede', 16), ('Qué se hizo', 50)], pares, 'Sin duplicados.')
    hoja_tabla(wb, 'Reemplazos', [('Archivo anterior', 30), ('Archivo que lo reemplaza', 30), ('Área', 30), ('Fila', 6),
                                  ('Actividad que no está en el nuevo', 50), ('Primera fecha', 13)],
               reemplazos, 'Ningún archivo nuevo dejó por fuera actividades del anterior.')
    hoja_tabla(wb, 'Archivos', [('Archivo', 34), ('Estado', 12), ('Motivo', 60), ('Área', 30), ('Unidad líder', 30),
                                ('Sede principal', 14), ('Quien diligencia', 24), ('Correo', 30), ('Extensión', 12),
                                ('Fecha de envío', 12), ('Visto bueno', 30), ('Filas escritas', 9),
                                ('Filas con error', 9), ('Fechas en el calendario', 10), ('Fecha del archivo', 12),
                                ('Huella', 14)], lista_archivos, 'No se encontraron archivos.')
    guardar(wb, REVISION)
    return resumen


# ── Programa ────────────────────────────────────────────────────────────
HOJAS_DE_CALCULO = {'.xlsx', '.xlsm', '.xls', '.xlsb', '.ods', '.csv'}
ENTRADA = HERRAMIENTAS / 'formatos'


def leer_carpeta(entrada: Path, anio: int | None = None) -> list[Archivo]:
    rutas = sorted(p for p in entrada.rglob('*') if p.is_file() and not p.name.startswith(('~$', '.')))
    return [leer_archivo(r, anio) if r.suffix.lower() in HOJAS_DE_CALCULO
            else Archivo(r, estado='Ignorado', motivo='No es una hoja de cálculo') for r in rutas]


def main() -> None:
    parser = argparse.ArgumentParser(description='Consolida los formatos de las áreas y los carga en la app.')
    parser.add_argument('--entrada', type=Path, default=ENTRADA, help='carpeta con los formatos (.xlsx)')
    parser.add_argument('--anio', type=int, help='avisa de las fechas que no son de este año')
    parser.add_argument('--solo-revisar', action='store_true', help='no reemplaza el Excel de la app: deja una vista previa')
    parser.add_argument('--reemplazar-todo', action='store_true',
                        help='borra también los otros años del Excel de la app (por defecto se conservan)')
    args = parser.parse_args()

    if not args.entrada.is_dir():
        args.entrada.mkdir(parents=True)
        sys.exit(f'Creé la carpeta «{args.entrada}»: ponga ahí los formatos de las áreas y vuelva a correr.')
    archivos = leer_carpeta(args.entrada, args.anio)
    if not any(a.estado != 'Ignorado' for a in archivos):
        sys.exit(f'No hay formatos en «{args.entrada}».')

    solo_en_anterior = resolver_reenvios(archivos)
    filas = [f for a in archivos if a.estado == 'Aceptado' for f in a.filas]
    filas, quitadas = quitar_repetidas(filas)
    duplicados = posibles_duplicados(filas)
    nuevas = filas_fuente(filas)
    origenes = [a.ruta.name for a in archivos if a.estado == 'Aceptado']

    # Los años que trae la consolidación se reemplazan; los demás años del Excel de la app se conservan.
    todas, conservadas = combinar_con_otros_anios(nuevas, conservar=not args.reemplazar_todo) if nuevas else ([], [])
    otros = Counter(f['desde'].year for f in conservadas)
    if otros:
        origenes.append('Filas de otros años del Excel anterior: ' + ', '.join(f'{a} ({n})' for a, n in sorted(otros.items())))
    cargado = False
    if args.solo_revisar:
        escribir_fuente(VISTA_PREVIA, todas or nuevas, 'herramientas/consolidar.py --solo-revisar', origenes)
        destino = [('Calendario', f'Vista previa, NO se cargó a la app: {VISTA_PREVIA.relative_to(RAIZ)}')]
    elif not nuevas:
        destino = [('Calendario', 'No hay actividades válidas: no se reemplazó el Excel de la app')]
    else:
        errores = cargar.validar(todas)
        if errores:  # no debería pasar: cada fila ya se revisó contra los mismos catálogos
            sys.exit('No se reemplazó nada: hay filas que la app no aceptaría.\n  ' + '\n  '.join(errores[:30]))
        antes = filas_en_fuente()
        respaldo = reemplazar_fuente(todas, 'herramientas/consolidar.py', origenes)
        cargado = True
        anios = ', '.join(map(str, sorted({f['desde'].year for f in nuevas})))
        destino = [('Excel de la app', f'Actualizado: {FUENTE.relative_to(RAIZ)} '
                                       f'({antes if antes is not None else "no existía"} → {len(todas)} filas; se reemplazó {anios})')]
        if otros:
            destino.append(('Otros años', 'Se conservaron sin cambios: ' + ', '.join(f'{a} ({n} filas)' for a, n in sorted(otros.items()))))
        if respaldo:
            destino.append(('Copia del anterior', str(respaldo.relative_to(RAIZ))))
        if antes and len(todas) < antes / 2:
            destino.append(('⚠ Atención', f'El calendario bajó de {antes} a {len(todas)} filas: ¿faltan formatos por '
                                          'llegar? Los años consolidados quedan solo con lo que llegó ahora.'))
    resumen = escribir_revision(args.entrada, archivos, filas, len(nuevas), quitadas, duplicados, solo_en_anterior,
                                args.anio, destino)

    for k, v in resumen.items():
        print(f'  {k}: {v}')
    for k, v in destino:
        print(f'{k}: {v}')
    print(f'Revisión: {REVISION}')
    if cargado:
        cargar.informar(cargar.cargar())


if __name__ == '__main__':
    main()
